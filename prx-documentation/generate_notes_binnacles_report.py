from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Informe_Tecnico_Notas_y_Bitacora_PRX.docx"

NAVY = "193753"
TEAL = "009B8E"
PALE = "EAF2F7"
LIGHT = "F4F6F8"
BORDER = "D9D9D9"
TEXT = RGBColor(35, 46, 58)


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    keep_with_next(p)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        r = p.add_run(f"{index}. ")
        r.bold = True
        p.add_run(item)


def add_label_paragraph(doc, label, text):
    p = doc.add_paragraph()
    r = p.add_run(label)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 128, 113)
    p.add_run(text)
    return p


def add_code(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.05)
    p.paragraph_format.right_indent = Inches(0.05)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F1F4F7")
    p_pr.append(shd)
    for line in code.strip().splitlines():
        run = p.add_run(line + "\n")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(30, 41, 59)
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(header))
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(font_size)
        if widths:
            cell.width = Inches(widths[i])
    for row_index, row in enumerate(rows):
        added_row = table.add_row()
        prevent_row_split(added_row)
        cells = added_row.cells
        for i, value in enumerate(row):
            cell = cells[i]
            if row_index % 2:
                set_cell_fill(cell, PALE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
            r.font.color.rgb = TEXT
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = TEXT
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08
for name, size in (("Title", 29), ("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12)):
    style = styles[name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
    style.paragraph_format.space_after = Pt(7)

title_p_pr = styles["Title"].element.get_or_add_pPr()
title_border = title_p_pr.find(qn("w:pBdr"))
if title_border is not None:
    title_p_pr.remove(title_border)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = header.add_run("PRX  Mantenimiento evolutivo")
run.font.name = "Aptos"
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(91, 105, 120)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Página ").font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

# Portada
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(48)
r = p.add_run("PROYECTO PRX")
r.bold = True
r.font.size = Pt(12)
r.font.color.rgb = RGBColor(0, 155, 142)

title = doc.add_paragraph(style="Title")
title.add_run("Informe técnico de Notas y Bitácora")

subtitle = doc.add_paragraph()
r = subtitle.add_run("Código implementado  calidad  SOLID  patrones y defensa")
r.bold = True
r.font.size = Pt(17)
r.font.color.rgb = RGBColor(25, 55, 83)

doc.add_paragraph(
    "Este documento explica el mantenimiento realizado en los módulos de Notas y Bitácora. "
    "Detalla las funciones añadidas, el recorrido de los datos desde la interfaz hasta MySQL, "
    "las decisiones de diseño y la forma de defender el trabajo según puntualidad, calidad, "
    "principios SOLID y patrones de diseño."
)

add_table(
    doc,
    ["Aspecto", "Contenido"],
    [
        ["Notas", "Enlaces URL, tareas, imágenes y editor de imágenes"],
        ["Bitácora", "Enlaces URL, tareas, vista separada de edición y guardado del checkbox"],
        ["Frontend", "Angular 21, Signals, formularios reactivos, PrimeIcons y Fabric.js"],
        ["Backend", "NestJS, CQRS, DTO, validación, permisos y repositorios"],
        ["Persistencia", "Prisma, MySQL, campos JSON, archivos privados y borrado lógico"],
        ["Verificación", "Build Angular y build NestJS ejecutados el 23 de septiembre de 2026"],
    ],
    [1.45, 5.65],
    9,
)

heading(doc, "Resultado principal", 2)
doc.add_paragraph(
    "Las notas dejaron de ser únicamente título y contenido: ahora relacionan recursos web, "
    "manejan tareas, adjuntan imágenes y permiten modificarlas antes de guardarlas. Las bitácoras "
    "también administran enlaces y tareas; su consulta está separada de la edición y el usuario "
    "puede completar una tarea directamente desde la vista."
)

doc.add_page_break()

heading(doc, "1 Alcance realizado", 1)
add_table(
    doc,
    ["Módulo", "Función", "Comportamiento entregado"],
    [
        ["Notas", "Enlaces URL", "Agregar, normalizar, validar, listar, abrir y eliminar enlaces."],
        ["Notas", "Tareas", "Agregar, eliminar, mostrar y marcar o desmarcar con persistencia."],
        ["Notas", "Imágenes", "Adjuntar JPEG, PNG y WebP, previsualizar, reemplazar y eliminar."],
        ["Notas", "Editor", "Seleccionar, dibujar, marcar, recortar, rotar, escribir, usar formas y deshacer."],
        ["Bitácora", "Enlaces URL", "Crear y editar listas de URL HTTP o HTTPS."],
        ["Bitácora", "Tareas", "Crear, eliminar y cambiar el estado desde la vista o edición."],
        ["Bitácora", "Detalle", "Abrir una página de consulta y entrar a Editar solo cuando se necesita."],
    ],
    [1.05, 1.35, 4.7],
)

heading(doc, "2 Tecnologías que usamos", 1)
add_table(
    doc,
    ["Tecnología", "Dónde", "Por qué se usó"],
    [
        ["Angular 21", "Frontend", "Componentes standalone, rutas, formularios y renderizado de las pantallas."],
        ["Signals", "Frontend", "Estado reactivo de tareas, enlaces, archivos, carga, edición y guardado."],
        ["Reactive Forms", "Frontend", "Validación de título y contenido antes de enviar datos."],
        ["Fabric.js", "Editor", "Canvas, dibujo, texto, figuras, selección, recorte e historial."],
        ["FormData", "Notas API", "Enviar en una misma petición texto, JSON y archivos binarios."],
        ["NestJS", "Backend", "Controladores, inyección de dependencias, validación y organización modular."],
        ["CQRS", "Backend", "Separar Commands que modifican datos de Queries que los consultan."],
        ["class validator", "DTO", "Rechazar tareas, enlaces y campos con formato o tamaño incorrectos."],
        ["Prisma", "Persistencia", "Acceso tipado a MySQL y mapeo de campos JSON."],
        ["Tigris Storage", "Archivos", "Guardar imágenes y adjuntos privados fuera de la base de datos."],
        ["JWT y Roles", "Seguridad", "Identificar al usuario y comprobar permisos antes de leer o modificar."],
    ],
    [1.35, 1.25, 4.5],
)

heading(doc, "3 Implementación de Notas", 1)
doc.add_paragraph(
    "La implementación de Notas se distribuyó entre presentación, contratos de dominio, API, casos de uso "
    "y persistencia. Esta separación evita que la pantalla tenga que conocer Prisma, MySQL o el almacenamiento."
)

heading(doc, "3.1 Modelos de tareas y enlaces", 2)
add_code(doc, """
export interface NoteTaskModel {
  id: string;
  title: string;
  completed: boolean;
}

export interface NoteLinkModel {
  id: string;
  url: string;
}
""")
add_label_paragraph(doc, "Qué hace  ", "Define la forma exacta de cada tarea y enlace en Angular.")
add_label_paragraph(doc, "Por qué se añadió  ", "Antes la nota no tenía contratos tipados para estas listas. El id permite localizar cada elemento, completed conserva el estado y url guarda la dirección normalizada.")

heading(doc, "3.2 Enlaces URL", 2)
add_bullets(doc, [
    "El usuario puede escribir una URL con o sin protocolo.",
    "normalizeUrl agrega https cuando falta el protocolo.",
    "isValidUrl usa la clase URL del navegador y restringe el resultado a HTTP o HTTPS.",
    "Cada enlace obtiene un identificador con crypto.randomUUID.",
    "La vista abre enlaces con target blank y rel noopener noreferrer para reducir riesgos.",
])
add_code(doc, r"""
private normalizeUrl(value: string): string {
  const trimmed = value.trim();
  if (!trimmed) return '';
  return /^https?:\/\//i.test(trimmed) ? trimmed : `https://${trimmed}`;
}

private isValidUrl(value: string): boolean {
  try {
    const url = new URL(value);
    return url.protocol === 'http:' || url.protocol === 'https:';
  } catch {
    return false;
  }
}
""")

heading(doc, "3.3 Tareas y guardado desde la vista", 2)
add_code(doc, """
protected toggleTask(task: NoteTaskModel): void {
  if (!this.canEdit() || this.saving()) return;
  this.tasks.update((tasks) => tasks.map((item) =>
    item.id === task.id ? { ...item, completed: !item.completed } : item
  ));
  if (!this.editing()) this.persist(false);
}
""")
doc.add_paragraph(
    "La lista se actualiza de forma inmutable. Cuando el usuario está en modo Ver, persist false envía el cambio "
    "sin abrir ni cerrar el editor. El bloqueo saving evita peticiones duplicadas mientras una actualización está en curso."
)

heading(doc, "3.4 Imágenes y archivos", 2)
add_bullets(doc, [
    "Se separan imágenes de otros adjuntos mediante el tipo MIME.",
    "Se aceptan JPEG, PNG y WebP y se comprueba el tamaño máximo.",
    "La nota mantiene un máximo total de cinco archivos entre existentes y nuevos.",
    "URL.createObjectURL crea una previsualización local y revokeObjectURL libera memoria.",
    "Una imagen ya guardada se descarga temporalmente, se convierte en File y puede volver al editor.",
    "Al guardar una edición se reemplaza el archivo original por una nueva imagen PNG.",
])

heading(doc, "3.5 Editor de imágenes", 2)
add_table(
    doc,
    ["Herramienta", "Código o API", "Resultado"],
    [
        ["Seleccionar", "Canvas selection", "Mover, redimensionar y seleccionar objetos."],
        ["Pincel", "PencilBrush", "Dibujar con color y grosor configurables."],
        ["Marcador", "PencilBrush con transparencia", "Resaltar zonas sin ocultar la imagen."],
        ["Recortar", "Rect y Canvas toDataURL", "Exportar únicamente el área seleccionada."],
        ["Texto", "IText", "Insertar y editar texto sobre la imagen."],
        ["Formas", "Rect Circle Triangle", "Añadir rectángulo, círculo y flecha."],
        ["Rotar", "Canvas 2D rotate", "Girar el resultado 90 grados."],
        ["Historial", "toJSON y loadFromJSON", "Deshacer y rehacer hasta 30 estados."],
        ["Guardar", "canvas toBlob", "Emitir un archivo PNG con nombre editada."],
    ],
    [1.25, 2.2, 3.65],
)
add_code(doc, """
const blob = await this.canvas.toBlob({ format: 'png', multiplier: 1 });
this.saved.emit(new File([blob], `${baseName}-editada.png`, {
  type: 'image/png',
  lastModified: Date.now(),
}));
""")

heading(doc, "3.6 Transporte al backend", 2)
add_code(doc, """
formData.append('tasks', JSON.stringify(data.tasks));
formData.append('links', JSON.stringify(data.links));
formData.append('retainedFileIds', data.retainedFileIds.join(','));
files.forEach((file) => formData.append('files', file, file.name));
""")
doc.add_paragraph(
    "Se usó FormData porque una actualización de nota puede combinar campos de texto, arreglos JSON y archivos. "
    "El backend vuelve a analizar las listas; no confía únicamente en la validación del navegador."
)

heading(doc, "3.7 Validación permisos y persistencia", 2)
add_bullets(doc, [
    "UpdateNoteHandler verifica que la nota y el repositorio existan.",
    "En repositorios íntimos solo el propietario puede modificar; en otros se acepta propietario o miembro.",
    "parseTasks limita a 100 tareas, exige id y título y recorta cada título a 200 caracteres.",
    "parseLinks limita a 50 enlaces, analiza cada URL y limita su longitud a 2048 caracteres.",
    "Los archivos eliminados usan soft delete y también se retiran del almacenamiento privado.",
    "Prisma guarda tasks y links como JSON y devuelve la nota completa mediante un Mapper.",
])

heading(doc, "4 Implementación de Bitácora", 1)
doc.add_paragraph(
    "Bitácora reutiliza la misma estructura conceptual de tareas y enlaces, pero transporta JSON normal en lugar "
    "de FormData porque no incluye archivos. Se añadió una página de detalle para separar claramente consulta y edición."
)

heading(doc, "4.1 Crear una bitácora con tareas y enlaces", 2)
add_code(doc, """
export interface CreateBinnacleRequest {
  name: string;
  content: string;
  tasks: { id: string; title: string; completed: boolean }[];
  links: { id: string; url: string }[];
}
""")
doc.add_paragraph(
    "El formulario de creación envía el título, el contenido y ambas listas en una sola petición. El Handler recorta "
    "espacios y construye BinnacleEntity antes de llamar al repositorio."
)

heading(doc, "4.2 Vista separada de edición", 2)
add_bullets(doc, [
    "La ruta de detalle obtiene el id y carga la bitácora con GET por id.",
    "En modo Ver se muestran contenido, fechas, enlaces y tareas; los campos de edición permanecen ocultos.",
    "El botón Editar cambia la señal editing y muestra controles para título, contenido, enlaces y tareas.",
    "Cancelar restablece los valores originales de la bitácora y descarta los cambios locales.",
    "Guardar valida el formulario, agrega entradas pendientes y ejecuta PUT por id.",
])

heading(doc, "4.3 Marcar tareas desde Ver", 2)
add_code(doc, """
protected toggleTask(id: string): void {
  if (this.saving()) return;
  this.tasks.update((tasks) =>
    tasks.map((task) => task.id === id
      ? { ...task, completed: !task.completed }
      : task),
  );
  if (!this.editing()) this.persist(false);
}
""")
doc.add_paragraph(
    "Este comportamiento cumple la solicitud de poder completar tareas sin entrar a Editar. La interfaz cambia el "
    "checkbox y llama a persist false. El backend comprueba que el userId del JWT sea el propietario antes de actualizar."
)

heading(doc, "4.4 Validaciones del backend", 2)
add_table(
    doc,
    ["Dato", "Regla", "Objetivo"],
    [
        ["Título", "2 a 15 caracteres", "Conservar el contrato existente de Bitácora."],
        ["Contenido", "1 a 2000 caracteres", "Evitar datos vacíos o excesivos."],
        ["Tareas", "Arreglo de máximo 50", "Controlar volumen y validar cada elemento."],
        ["Título de tarea", "No vacío y máximo 250", "Mantener tareas legibles."],
        ["Enlaces", "Arreglo de máximo 50", "Controlar volumen de referencias."],
        ["URL", "HTTP o HTTPS con protocolo", "Rechazar esquemas y formatos no esperados."],
        ["Permiso", "Rol estándar y propietario", "Impedir editar bitácoras ajenas."],
    ],
    [1.35, 2.2, 3.55],
)

heading(doc, "4.5 Código del caso de uso", 2)
add_code(doc, """
if (binnacle.userId !== command.userId) {
  throw new ForbiddenException(BINNACLE_MESSAGES.FORBIDDEN);
}

const updated = await this.binnacleRepository.update(command.id, {
  name: command.dto.name.trim(),
  content: command.dto.content.trim(),
  tasks: command.dto.tasks.map((task) => ({ ...task, title: task.title.trim() })),
  links: command.dto.links.map((link) => ({ ...link, url: link.url.trim() })),
  updatedBy: command.userId,
});
""")

heading(doc, "5 Base de datos y recorrido de los datos", 1)
heading(doc, "5.1 Campos agregados", 2)
add_code(doc, """
model Note {
  tasks Json?
  links Json?
}

model Binnacle {
  tasks Json?
  links Json?
}
""")
doc.add_paragraph(
    "Sí existen campos para guardar tareas y enlaces. Se usaron columnas JSON opcionales porque cada lista pertenece "
    "a una sola nota o bitácora y su estructura es pequeña. Prisma convierte los arreglos a InputJsonValue al crear o actualizar."
)

heading(doc, "5.2 Flujo completo", 2)
add_numbered(doc, [
    "El usuario agrega una tarea, URL o imagen en Angular.",
    "El componente valida y actualiza Signals sin mutar el arreglo anterior.",
    "La Facade coordina la llamada y la API construye JSON o FormData.",
    "El Controller extrae el usuario del JWT y envía un Command al CommandBus.",
    "El DTO valida formato y límites; el Handler valida permisos y reglas del caso de uso.",
    "El Repository usa Prisma para guardar tasks y links como JSON en MySQL.",
    "El Mapper transforma el registro de Prisma en entidad y luego en DTO de respuesta.",
    "El Store y el componente actualizan la interfaz con el resultado confirmado por el servidor.",
])

heading(doc, "6 Código incorporado por capa", 1)
add_table(
    doc,
    ["Capa", "Archivos principales", "Responsabilidad añadida"],
    [
        ["Presentación Notas", "create note page  note detail page", "Controles URL, tareas, imágenes, vista y edición."],
        ["Editor Notas", "note image editor component", "Canvas Fabric.js y herramientas gráficas."],
        ["Dominio Notas", "note model  create y update request", "Contratos de Task, Link y archivos retenidos."],
        ["API Notas", "note api", "FormData con tasks, links y archivos."],
        ["Backend Notas", "DTO  Handler  Entity", "Parseo, permisos, reglas y normalización."],
        ["Persistencia Notas", "Mapper  PrismaNoteRepository", "Lectura y escritura de JSON."],
        ["Presentación Bitácora", "binnacle page  detail page", "Crear, ver, editar y marcar tareas."],
        ["Dominio Bitácora", "model  requests  contracts", "Contratos de Task y Link."],
        ["Backend Bitácora", "DTO  Controllers  Handlers", "Rutas por id, validación y propiedad."],
        ["Persistencia Bitácora", "Mapper  PrismaBinnacleRepository", "Guardar y recuperar JSON."],
        ["Esquema", "prisma schema", "Campos tasks y links en Note y Binnacle."],
    ],
    [1.2, 2.5, 3.4],
    8.2,
)

heading(doc, "7 Calidad del trabajo", 1)
heading(doc, "7.1 Evidencia comprobada", 2)
add_table(
    doc,
    ["Comprobación", "Resultado", "Lectura para la defensa"],
    [
        ["Build Angular", "Correcto", "La aplicación frontend genera el bundle sin errores."],
        ["Build NestJS", "Correcto", "Prisma Client se genera y TypeScript compila sin errores."],
        ["Validación URL", "Frontend y backend", "No se depende de una sola capa para validar datos."],
        ["Permisos", "Propiedad o membresía", "Los cambios no exponen datos de otros usuarios."],
        ["Persistencia", "Campos JSON", "Los enlaces y estados de tareas permanecen tras recargar."],
        ["Archivos", "Límite y tipos", "Máximo cinco archivos y formatos de imagen definidos."],
        ["Guardado", "Estados saving", "Se evita repetir acciones mientras una petición está activa."],
    ],
    [1.45, 1.65, 4.0],
)

heading(doc, "7.2 Detalles que demuestran calidad", 2)
add_bullets(doc, [
    "Los inputs tienen límites coherentes con los DTO y el esquema de base de datos.",
    "Los errores HTTP se transforman en notificaciones comprensibles para el usuario.",
    "Los enlaces externos usan noopener noreferrer.",
    "El editor libera URL temporales y destruye Canvas y ResizeObserver al cerrar.",
    "Los Mappers devuelven arreglos vacíos cuando un campo JSON es nulo o antiguo.",
    "El borrado es lógico mediante status igual a cero para conservar auditoría.",
])

heading(doc, "7.3 Aspectos pendientes", 2)
add_bullets(doc, [
    "No se encontraron pruebas automatizadas específicas para las nuevas funciones; deben añadirse pruebas unitarias y de integración.",
    "El build Angular muestra una advertencia: create note page component SCSS supera el presupuesto de 5 KB por 818 bytes.",
    "Conviene probar archivos grandes, fallos de red y edición simultánea antes de producción.",
])
doc.add_paragraph(
    "Mencionar estas mejoras no invalida el Sprint. Demuestra que se conoce el estado real del producto y que existe un plan de calidad continuo."
)

heading(doc, "8 Principios SOLID aplicados", 1)
add_table(
    doc,
    ["Principio", "Dónde se observa", "Por qué es útil"],
    [
        ["S Responsabilidad única", "Component  Facade  API  Handler  Repository  Mapper", "Cada clase se concentra en interfaz, coordinación, HTTP, negocio, datos o conversión."],
        ["O Abierto cerrado", "Contratos y componentes aislados", "Se amplían capacidades sin reescribir todas las capas consumidoras."],
        ["L Sustitución de Liskov", "PrismaNoteRepository implementa NoteRepository", "La aplicación usa el contrato sin depender del detalle de Prisma."],
        ["I Segregación de interfaces", "NoteApiContract y BinnacleApiContract separados", "Cada módulo depende solo de las operaciones que necesita."],
        ["D Inversión de dependencias", "Handlers inyectan Repository abstracto", "El caso de uso no crea ni conoce directamente PrismaService."],
    ],
    [1.45, 2.45, 3.2],
    8.2,
)

heading(doc, "Cómo defender SOLID", 2)
doc.add_paragraph(
    "No colocamos todo el código en el componente. La pantalla administra interacción; la Facade coordina; la API transporta; "
    "el DTO y el Handler validan; el Repository persiste; el Mapper traduce formatos. Si mañana cambiamos Prisma por otra "
    "tecnología, el Handler puede seguir usando el contrato del repositorio. Ese es el ejemplo más claro de SRP y DIP."
)

heading(doc, "9 Patrones de diseño utilizados", 1)
add_table(
    doc,
    ["Patrón", "Uso real", "Por qué se utilizó"],
    [
        ["CQRS", "Commands para crear o actualizar y Queries para consultar", "Separar lectura y escritura y mantener casos de uso pequeños."],
        ["Repository", "NoteRepository y BinnacleRepository", "Aislar reglas de negocio de Prisma y MySQL."],
        ["Facade", "NoteFacade y BinnacleFacade", "Dar a la pantalla una API simple y coordinar API y Store."],
        ["Mapper", "Prisma Mapper y Response Mapper", "Evitar que formatos de base de datos se filtren a la UI."],
        ["DTO", "Create y Update DTO", "Definir y validar la entrada de cada endpoint."],
        ["Reactive Store", "Signals y Stores", "Mantener estado predecible de carga, datos y errores."],
        ["Guard y Decorator", "JWT CurrentUser y Roles", "Centralizar autenticación y autorización."],
        ["Soft delete", "status igual a cero", "Conservar trazabilidad y auditoría sin borrado físico inmediato."],
    ],
    [1.3, 2.75, 3.05],
    8.2,
)

heading(doc, "Por qué no son patrones decorativos", 2)
doc.add_paragraph(
    "Cada patrón resuelve un problema presente en el mantenimiento. CQRS organiza las operaciones; Repository desacopla "
    "persistencia; Facade reduce dependencias de la pantalla; Mapper protege los límites entre capas; DTO evita datos inválidos; "
    "los Guards impiden acceso no autorizado. No se añadieron solo para nombrarlos en la exposición."
)

heading(doc, "10 Puntualidad", 1)
doc.add_paragraph(
    "La entrega indicada vence el 23 de septiembre de 2026 a las 23:59. La puntualidad no se demuestra únicamente diciendo "
    "que se trabajó a tiempo: debe respaldarse con el archivo entregado, el registro de envío y una demostración preparada antes del cierre."
)
add_table(
    doc,
    ["Evidencia", "Estado", "Acción"],
    [
        ["Funciones principales", "Implementadas", "Demostrar Notas y Bitácora de principio a fin."],
        ["Compilación frontend", "Correcta", "Conservar la salida del build como evidencia."],
        ["Compilación backend", "Correcta", "Conservar la salida de Prisma y Nest build."],
        ["Documento de defensa", "Preparado", "Adjuntar este informe antes de las 23:59."],
        ["Envío en plataforma", "Responsabilidad del estudiante", "Confirmar el envío y guardar captura con fecha y hora."],
    ],
    [1.7, 1.55, 3.85],
)

heading(doc, "11 Guion para presentar y defender", 1)
add_table(
    doc,
    ["Tiempo", "Tema", "Qué mostrar o explicar"],
    [
        ["1 min", "Alcance", "Notas con URL, tareas, imágenes y editor; Bitácora con URL y tareas."],
        ["2 min", "Notas", "Agregar URL y tarea, editar imagen, guardar, recargar y marcar tarea."],
        ["2 min", "Bitácora", "Crear, abrir Ver, marcar tarea, entrar a Editar y guardar."],
        ["2 min", "Código", "Modelos, FormData o JSON, DTO, Handler, Repository y schema Prisma."],
        ["2 min", "Calidad", "Validaciones, permisos, límites, build correcto y mejoras pendientes."],
        ["2 min", "SOLID y patrones", "SRP, DIP, CQRS, Repository, Facade, Mapper y DTO."],
        ["1 min", "Cierre", "Valor entregado y estado de la solución."],
    ],
    [0.8, 1.35, 4.95],
)

heading(doc, "11.1 Demostración paso a paso", 2)
add_numbered(doc, [
    "Iniciar sesión y abrir una nota existente.",
    "Entrar a Editar, agregar una URL sin protocolo y comprobar que se normaliza.",
    "Agregar una tarea y seleccionar una imagen JPEG, PNG o WebP.",
    "Abrir el editor, dibujar o agregar texto, deshacer y guardar la imagen editada.",
    "Guardar la nota, recargar y comprobar enlaces, archivos y tareas.",
    "Marcar una tarea desde Ver sin entrar de nuevo a Editar.",
    "Abrir Bitácora, crear un registro con URL y tareas.",
    "Entrar al detalle, marcar una tarea desde Ver y recargar.",
    "Mostrar que Ver y Editar son estados separados.",
    "Cerrar enseñando schema Prisma, Handler y Repository.",
])

heading(doc, "11.2 Mensaje de apertura", 2)
doc.add_paragraph(
    "El mantenimiento asignado amplió dos módulos existentes. En Notas incorporamos enlaces URL, tareas, imágenes y un "
    "editor gráfico. En Bitácora agregamos enlaces y tareas, además de separar la consulta de la edición. La solución no se "
    "limitó a la interfaz: incluye validación, permisos, casos de uso y persistencia en MySQL."
)

heading(doc, "11.3 Mensaje de cierre", 2)
doc.add_paragraph(
    "El resultado permite relacionar información externa, seguir actividades y trabajar visualmente con imágenes sin salir de "
    "la nota. En Bitácora, el usuario consulta primero y edita solo cuando lo necesita, pero puede completar tareas de forma directa. "
    "La arquitectura mantiene las responsabilidades separadas, protege los datos y permite continuar mejorando cada módulo."
)

heading(doc, "12 Preguntas probables y respuestas", 1)
questions = [
    ("¿Dónde se guardan las tareas y los enlaces?", "En los campos JSON opcionales tasks y links de Note y Binnacle en MySQL, accedidos mediante Prisma."),
    ("¿Por qué JSON y no tablas nuevas?", "Son listas pequeñas, dependientes de un solo registro y se recuperan siempre junto con él. JSON redujo complejidad para este alcance; si luego requieren consultas globales o relaciones propias convendría normalizarlas."),
    ("¿Cómo validan una URL?", "Angular normaliza y usa new URL; el backend vuelve a validar. En Bitácora class validator exige HTTP o HTTPS con protocolo."),
    ("¿Cómo evitan editar datos ajenos?", "El JWT entrega el userId. Bitácora compara ese id con el propietario; Notas valida propietario o membresía del repositorio según su visibilidad."),
    ("¿Qué pasa al marcar una tarea desde Ver?", "La señal actualiza la lista y persist false ejecuta PUT sin activar el editor. El servidor devuelve el registro actualizado."),
    ("¿Por qué Fabric.js?", "Resuelve selección, objetos, dibujo, serialización e interacción sobre Canvas, evitando implementar un motor gráfico manual."),
    ("¿Dónde está SOLID?", "En la separación Component, Facade, API, Handler, Repository y Mapper; además los Handlers dependen de contratos de repositorio."),
    ("¿Qué patrones usaron?", "CQRS, Repository, Facade, Mapper, DTO, Store reactivo, Guards y soft delete."),
    ("¿Cómo demostraron calidad?", "Con validaciones en dos capas, permisos, límites, manejo de errores y compilaciones correctas de Angular y NestJS."),
    ("¿Qué falta mejorar?", "Pruebas automatizadas específicas, reducción del SCSS que supera el presupuesto y pruebas adicionales de archivos grandes y fallos de red."),
]
for question, answer in questions:
    p = doc.add_paragraph()
    r = p.add_run(question)
    r.bold = True
    r.font.color.rgb = RGBColor(25, 55, 83)
    p = doc.add_paragraph(answer)
    p.paragraph_format.space_after = Pt(8)

heading(doc, "13 Conclusión", 1)
doc.add_paragraph(
    "El mantenimiento cubre el alcance funcional solicitado en Notas y Bitácora y lo integra con la arquitectura existente. "
    "Las mejoras tienen contratos tipados, validación, autorización, persistencia y respuesta visual. Las compilaciones pasan; "
    "el siguiente paso de calidad es incorporar pruebas automatizadas y atender la advertencia de tamaño de estilos."
)

heading(doc, "Anexo de archivos de referencia", 1)
add_bullets(doc, [
    "prx-frontend/src/app/features/notes",
    "prx-frontend/src/app/features/binnacles",
    "prx-backend/src/modules/notes",
    "prx-backend/src/modules/binnacles",
    "prx-backend/prisma/schema.prisma",
])

doc.core_properties.title = "Informe técnico de Notas y Bitácora"
doc.core_properties.subject = "Mantenimiento evolutivo PRX y defensa"
doc.core_properties.author = "Proyecto PRX"
doc.save(OUTPUT)
print(OUTPUT)
