from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Backlog_y_Defensa_Sprint_1_Mantenimiento_PRX.docx"

NAVY = "17324D"
TEAL = "009F92"
PALE_BLUE = "EEF4F8"
GRAY = "D9E0E5"
DARK = "202A33"
MUTED = "5C6975"
GREEN = "16794B"
AMBER = "9A6700"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_format(cell, fill=None):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), GRAY)
        borders.append(node)
    margins = OxmlElement("w:tcMar")
    for name, value in (("top", 95), ("start", 110), ("bottom", 95), ("end", 110)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)
    if fill:
        shade(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def cant_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def keep(paragraph):
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    keep(p)
    return p


def bullet(doc, text, level=0, size=10):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p


def numbered(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{number}. ")
    r.bold = True
    p.add_run(text)
    return p


def status(doc, label, color, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(f"{label}: ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(color)
    p.add_run(text)


def table(doc, headers, rows, widths, font_size=8.2):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    repeat_header(t.rows[0])
    for index, title in enumerate(headers):
        c = t.rows[0].cells[index]
        cell_format(c, NAVY)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(font_size)
        r.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        row = t.add_row()
        cant_split(row)
        for index, value in enumerate(values):
            c = row.cells[index]
            cell_format(c, "FFFFFF" if row_index % 2 == 0 else PALE_BLUE)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in (0, len(values) - 1) else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
            r.font.color.rgb = RGBColor.from_string(DARK)
    for row in t.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def code(doc, caption, source):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(caption)
    r.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    keep(p)
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    t.columns[0].width = Cm(16.2)
    cant_split(t.rows[0])
    c = t.cell(0, 0)
    cell_format(c, "F3F5F7")
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(source.strip())
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    r.font.size = Pt(7.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2.1)
section.right_margin = Cm(2.1)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.2)
normal.font.color.rgb = RGBColor.from_string(DARK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

title_style = doc.styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(30)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_ppr = title_style._element.get_or_add_pPr()
border = title_ppr.find(qn("w:pBdr"))
if border is not None:
    title_ppr.remove(border)

for style_name, size in (("Heading 1", 19), ("Heading 2", 14), ("Heading 3", 11.5)):
    style = doc.styles[style_name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(14 if style_name == "Heading 1" else 10)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.text = "PRX  |  Mantenimiento evolutivo"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
page_number(section.footer.paragraphs[0])

# Portada
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(40)
r = p.add_run("PROYECTO PRX")
r.bold = True
r.font.size = Pt(11)
r.font.color.rgb = RGBColor.from_string(TEAL)

p = doc.add_paragraph(style="Title")
p.add_run("Backlog de mantenimiento y defensa del Sprint 1")

p = doc.add_paragraph()
r = p.add_run("Notas y Bitácora")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = RGBColor.from_string(NAVY)
p.paragraph_format.space_after = Pt(22)

doc.add_paragraph(
    "Documento de planificación y defensa del mantenimiento evolutivo asignado al proyecto PRX. El trabajo funcional solicitado se limita a ampliar Notas con enlaces URL, tareas, imágenes y edición de imágenes, y ampliar Bitácora con enlaces URL y tareas."
)

table(
    doc,
    ["Dato", "Detalle"],
    [
        ["Entrega", "Presentación y defensa del Sprint 1"],
        ["Vencimiento", "23 de septiembre de 2026 a las 23:59"],
        ["Alcance funcional", "Notas y Bitácora"],
        ["Organización", "Tres sprints incrementales"],
        ["Tecnologías", "Angular 21, PrimeNG, Fabric.js, NestJS, CQRS, Prisma y MySQL"],
    ],
    [4.0, 12.2],
    9,
)

heading(doc, "Resultado esperado", 2)
doc.add_paragraph(
    "Al finalizar los tres sprints, las notas podrán relacionar recursos mediante URL, manejar listas de tareas, adjuntar imágenes y modificarlas con herramientas gráficas. Las bitácoras podrán guardar enlaces y tareas, mostrar el detalle separado de la edición y persistir el marcado de tareas."
)

doc.add_page_break()

# Resumen
heading(doc, "1 Alcance correcto del mantenimiento", 1)
doc.add_paragraph(
    "Este mantenimiento no tiene como objetivo construir un Home, buscador general u otros módulos del producto. El backlog se deriva únicamente de las mejoras solicitadas para Notas y Bitácora. La autenticación se documenta como dependencia técnica porque aporta el usuario y los permisos, pero no constituye una historia funcional adicional del mantenimiento."
)

table(
    doc,
    ["Módulo", "Función solicitada", "Resultado funcional"],
    [
        ["Notas", "Enlaces URL", "Agregar, validar, mostrar, abrir y eliminar enlaces relacionados."],
        ["Notas", "Tareas", "Crear, eliminar, visualizar y marcar tareas como completadas."],
        ["Notas", "Imágenes", "Adjuntar, previsualizar, quitar y guardar imágenes JPEG, PNG y WebP."],
        ["Notas", "Editor de imágenes", "Seleccionar, dibujar, marcar, recortar, rotar, añadir texto y formas, deshacer y rehacer."],
        ["Bitácora", "Enlaces URL", "Agregar, validar, mostrar, abrir y eliminar enlaces."],
        ["Bitácora", "Tareas", "Crear, eliminar, visualizar y marcar tareas con guardado inmediato."],
    ],
    [3.0, 4.0, 9.2],
)

heading(doc, "2 Product Backlog dividido en tres sprints", 1)
doc.add_paragraph(
    "La división minimiza dependencias: primero se establece el modelo de listas JSON y la interacción en Notas; después se incorpora el procesamiento gráfico; finalmente se reutilizan las decisiones aprendidas en Bitácora."
)

table(
    doc,
    ["ID", "Historia de usuario", "Prioridad", "Puntos", "Sprint"],
    [
        ["HU M01", "Como usuario quiero agregar enlaces URL a una nota para relacionarla con recursos externos.", "Alta", "5", "1"],
        ["HU M02", "Como usuario quiero administrar tareas en una nota para controlar actividades pendientes y completadas.", "Alta", "8", "1"],
        ["HU M03", "Como usuario quiero adjuntar imágenes a una nota para complementar visualmente su contenido.", "Alta", "5", "2"],
        ["HU M04", "Como usuario quiero editar una imagen antes de guardarla para resaltar o corregir información.", "Alta", "13", "2"],
        ["HU M05", "Como usuario quiero agregar enlaces URL a una bitácora para conservar referencias de mi trabajo.", "Media", "5", "3"],
        ["HU M06", "Como usuario quiero administrar tareas en una bitácora para registrar el avance diario.", "Alta", "8", "3"],
    ],
    [1.8, 8.8, 2.0, 1.5, 1.6],
    7.7,
)

heading(doc, "Resumen de sprints", 2)
table(
    doc,
    ["Sprint", "Objetivo", "Historias", "Puntos", "Entregable"],
    [
        ["Sprint 1", "Agregar enlaces URL y tareas en Notas.", "HU M01 y HU M02", "13", "Notas relacionadas y accionables."],
        ["Sprint 2", "Agregar imágenes y herramientas de edición en Notas.", "HU M03 y HU M04", "18", "Flujo visual con editor gráfico."],
        ["Sprint 3", "Agregar enlaces URL y tareas en Bitácora.", "HU M05 y HU M06", "13", "Bitácora enriquecida y editable."],
    ],
    [2.0, 6.0, 3.0, 1.5, 3.7],
    8,
)

# Sprint 1
heading(doc, "3 Sprint 1 Enlaces URL y tareas en Notas", 1)
status(doc, "Objetivo", GREEN, "permitir que una nota contenga recursos web y una lista de actividades que pueda actualizarse desde la vista.")
status(doc, "Valor", TEAL, "la nota deja de ser texto aislado y se convierte en un elemento de seguimiento y consulta.")
status(doc, "Vencimiento", AMBER, "23 de septiembre de 2026 a las 23:59.")

heading(doc, "3.1 Historia HU M01 Enlaces URL", 2)
doc.add_paragraph(
    "Como usuario quiero agregar enlaces URL a una nota para acceder desde ella a documentación, sitios o recursos relacionados."
)
heading(doc, "Criterios de aceptación", 3)
for item in [
    "El usuario puede escribir una dirección con o sin protocolo; el sistema normaliza el valor.",
    "Solo se aceptan direcciones HTTP o HTTPS válidas.",
    "Se pueden agregar varios enlaces y eliminar cualquiera antes de guardar.",
    "Los enlaces guardados se muestran en el detalle de la nota.",
    "Cada enlace se abre en una pestaña nueva con protección noopener noreferrer.",
    "Los enlaces permanecen después de recargar la página.",
]:
    bullet(doc, item)

heading(doc, "3.2 Historia HU M02 Tareas", 2)
doc.add_paragraph(
    "Como usuario quiero agregar y completar tareas dentro de una nota para registrar las actividades asociadas a su contenido."
)
heading(doc, "Criterios de aceptación", 3)
for item in [
    "El usuario puede agregar una tarea con título no vacío.",
    "Puede eliminar tareas durante la edición.",
    "El detalle muestra tareas pendientes y completadas.",
    "El checkbox permite marcar o desmarcar una tarea sin entrar al modo de edición.",
    "El cambio se guarda inmediatamente mediante el endpoint de actualización.",
    "Las tareas permanecen después de recargar la página.",
]:
    bullet(doc, item)

heading(doc, "3.3 Sprint Backlog", 2)
table(
    doc,
    ["Tarea", "Trabajo técnico", "Ubicación", "Estado"],
    [
        ["T1", "Definir NoteTaskModel y NoteLinkModel.", "domain/models/note.model.ts", "Hecho"],
        ["T2", "Crear controles para agregar y quitar URL.", "create-note y note-detail HTML/TS", "Hecho"],
        ["T3", "Normalizar y validar HTTP/HTTPS.", "create-note y note-detail TS", "Hecho"],
        ["T4", "Crear controles de tareas y checkbox.", "note-detail HTML/TS", "Hecho"],
        ["T5", "Enviar links y tasks mediante FormData.", "infrastructure/api/note.api.ts", "Hecho"],
        ["T6", "Validar y transformar JSON en backend.", "UpdateNoteHandler", "Hecho"],
        ["T7", "Persistir tasks y links como JSON.", "schema.prisma y PrismaNoteRepository", "Hecho"],
        ["T8", "Aplicar permisos de propietario o miembro.", "Handlers de Notes", "Hecho"],
        ["T9", "Probar guardado, recarga y URLs inválidas.", "Prueba funcional", "Verificado"],
        ["T10", "Preparar defensa y evidencia técnica.", "Documento y demostración", "Listo"],
    ],
    [1.4, 6.0, 6.0, 2.0],
    7.7,
)

heading(doc, "3.4 Definición de terminado", 2)
for item in [
    "La interfaz permite administrar enlaces y tareas sin errores de maquetación.",
    "Frontend y backend compilan.",
    "Los datos se guardan y se recuperan desde MySQL mediante Prisma.",
    "Las entradas inválidas reciben mensajes comprensibles.",
    "Los permisos impiden modificar notas sin acceso al repositorio.",
    "La demostración completa puede repetirse antes de la entrega.",
]:
    bullet(doc, item)

# Evidencia técnica Sprint1
heading(doc, "4 Evidencia técnica del Sprint 1", 1)
heading(doc, "4.1 Modelo de dominio en Angular", 2)
code(
    doc,
    "Archivo: prx-frontend/src/app/features/notes/domain/models/note.model.ts",
    """export interface NoteTaskModel {
  id: string;
  title: string;
  completed: boolean;
}

export interface NoteLinkModel {
  id: string;
  url: string;
}""",
)
doc.add_paragraph(
    "Estas interfaces tipan los elementos antes de enviarlos al backend. El id permite localizar un elemento, title describe la tarea, completed conserva su estado y url contiene la dirección normalizada."
)

heading(doc, "4.2 Validación de enlaces", 2)
code(
    doc,
    "Archivos: create-note-page.component.ts y note-detail-page.component.ts",
    """private isValidUrl(value: string): boolean {
  try {
    const url = new URL(value);
    return url.protocol === 'http:' || url.protocol === 'https:';
  } catch {
    return false;
  }
}""",
)
doc.add_paragraph(
    "La validación utiliza URL del navegador en lugar de expresiones regulares improvisadas. Además de comprobar la estructura, limita el protocolo para evitar valores no esperados."
)

heading(doc, "4.3 Persistencia inmediata del checkbox", 2)
code(
    doc,
    "Archivo: note-detail-page.component.ts",
    """protected toggleTask(task: NoteTaskModel): void {
  this.tasks.update((tasks) =>
    tasks.map((item) => item.id === task.id
      ? { ...item, completed: !item.completed }
      : item),
  );
  if (!this.editing()) this.persist(false);
}""",
)
doc.add_paragraph(
    "La lista se actualiza de manera inmutable. Si el usuario está consultando la nota, persist(false) guarda el cambio sin abrir el editor, ofreciendo una interacción rápida."
)

heading(doc, "4.4 Transporte y almacenamiento", 2)
code(
    doc,
    "Archivo: prx-frontend/src/app/features/notes/infrastructure/api/note.api.ts",
    """formData.append('tasks', JSON.stringify(data.tasks));
formData.append('links', JSON.stringify(data.links));""",
)
code(
    doc,
    "Archivo: prx-backend/prisma/schema.prisma",
    """model Note {
  title   String
  content String
  tasks   Json?
  links   Json?
}""",
)
doc.add_paragraph(
    "Las listas se serializan porque la misma solicitud también transporta archivos. El backend vuelve a analizar y validar el JSON antes de guardarlo. Prisma utiliza campos Json opcionales porque las listas pertenecen a una nota y se recuperan junto con ella."
)

heading(doc, "4.5 Archivos principales modificados", 2)
table(
    doc,
    ["Capa", "Archivo", "Responsabilidad"],
    [
        ["Presentación", "notes-page.component.*", "Tarjetas, vista previa y navegación."],
        ["Presentación", "create-note-page.component.*", "Creación y enlaces URL."],
        ["Presentación", "note-detail-page.component.*", "Vista, edición, enlaces y tareas."],
        ["Aplicación", "note.facade.ts", "Coordinar API, loading y Store."],
        ["Dominio", "note.model.ts / requests", "Contratos tipados."],
        ["Infraestructura", "note.api.ts / note.store.ts", "HTTP, FormData y estado."],
        ["Backend", "notes.controller.ts", "Endpoints GET, POST, PUT y DELETE."],
        ["Backend", "UpdateNoteHandler", "Permisos, validación y caso de uso."],
        ["Persistencia", "schema.prisma / PrismaNoteRepository", "Campos JSON y acceso MySQL."],
    ],
    [3.0, 6.2, 7.0],
    7.8,
)

# Sprint 2
heading(doc, "5 Sprint 2 Imágenes y editor en Notas", 1)
status(doc, "Objetivo", GREEN, "permitir adjuntar imágenes y modificarlas antes de guardarlas en la nota.")

heading(doc, "5.1 Historias y aceptación", 2)
table(
    doc,
    ["Historia", "Criterios de aceptación principales"],
    [
        ["HU M03 Imágenes", "Aceptar JPEG, PNG y WebP; mostrar vista previa; permitir quitar; respetar el máximo conjunto de cinco archivos; guardar y volver a visualizar."],
        ["HU M04 Editor", "Abrir cada imagen en un editor; aplicar herramientas; guardar como PNG; reemplazar la vista previa; conservar deshacer y rehacer."],
    ],
    [5.0, 11.2],
)

heading(doc, "5.2 Herramientas implementadas", 2)
table(
    doc,
    ["Herramienta", "Qué hace"],
    [
        ["Seleccionar", "Mover, redimensionar o seleccionar objetos agregados."],
        ["Pincel", "Dibujar trazos libres con color y grosor configurables."],
        ["Marcar", "Aplicar resaltado semitransparente."],
        ["Recortar", "Definir un área y reemplazar el lienzo con el recorte."],
        ["Texto", "Agregar texto editable sobre la imagen."],
        ["Rotar 90 grados", "Girar todo el resultado un cuarto de vuelta."],
        ["Formas", "Agregar rectángulos, círculos y flechas."],
        ["Eliminar selección", "Quitar el objeto gráfico activo."],
        ["Deshacer y rehacer", "Recorrer el historial de cambios del lienzo."],
        ["Guardar cambios", "Convertir el lienzo final en un archivo PNG."],
    ],
    [5.0, 11.2],
    8,
)

heading(doc, "5.3 Tareas del sprint", 2)
for item in [
    "Separar adjuntos generales de imágenes editables.",
    "Validar formatos JPEG, PNG y WebP y el máximo total de archivos.",
    "Crear previsualización con URL.createObjectURL y liberar recursos al eliminar.",
    "Integrar Fabric.js y un canvas adaptable.",
    "Implementar modos seleccionar, pincel, marcador y recorte.",
    "Agregar texto, rectángulo, círculo y flecha.",
    "Implementar rotación, historial, deshacer y rehacer.",
    "Exportar el resultado a PNG y sustituir la imagen editada.",
    "Permitir editar también una imagen ya guardada en una nota.",
]:
    bullet(doc, item)

# Sprint 3
heading(doc, "6 Sprint 3 Enlaces URL y tareas en Bitácora", 1)
status(doc, "Objetivo", GREEN, "enriquecer cada registro diario con referencias web y una lista de tareas persistente.")

heading(doc, "6.1 Historias y aceptación", 2)
table(
    doc,
    ["Historia", "Criterios de aceptación principales"],
    [
        ["HU M05 Enlaces", "Agregar varias URL HTTP/HTTPS, eliminar, guardar, listar y abrir de forma segura."],
        ["HU M06 Tareas", "Agregar, eliminar, marcar, desmarcar y conservar tareas después de recargar."],
    ],
    [5.0, 11.2],
)

heading(doc, "6.2 Tareas del sprint", 2)
for item in [
    "Añadir BinnacleTaskModel y BinnacleLinkModel.",
    "Incorporar controles de tareas y URL al formulario de creación.",
    "Agregar campos tasks y links al modelo Prisma Binnacle.",
    "Actualizar DTO, entidad, mapper y repositorio Prisma.",
    "Implementar GET por ID y PUT por ID.",
    "Separar la acción Ver de la acción Editar.",
    "Permitir marcar tareas desde la vista y guardar inmediatamente.",
    "Restringir el módulo al rol estandar y a registros propios.",
]:
    bullet(doc, item)

heading(doc, "6.3 Dependencias de seguridad", 2)
doc.add_paragraph(
    "Bitácora utiliza @Roles(Role.estandar) y compara userId con el sub del JWT. Notas comprueba propiedad o membresía del repositorio. Estas reglas pertenecen al soporte de autenticación existente y permiten demostrar que las nuevas funciones no exponen información de otros usuarios."
)

# SOLID and patterns
heading(doc, "7 SOLID aplicado al mantenimiento", 1)
table(
    doc,
    ["Principio", "Aplicación concreta"],
    [
        ["S Responsabilidad única", "El componente maneja interacción; la Facade coordina; el Handler aplica reglas; el repositorio persiste."],
        ["O Abierto cerrado", "Los contratos permiten ampliar implementaciones sin cambiar las páginas consumidoras."],
        ["L Sustitución de Liskov", "Las implementaciones Prisma respetan los contratos NoteRepository y BinnacleRepository."],
        ["I Segregación de interfaces", "Notas, Bitácora y Auth tienen contratos separados y enfocados."],
        ["D Inversión de dependencias", "Los Handlers dependen de repositorios abstractos inyectados, no de Prisma directamente."],
    ],
    [4.0, 12.2],
)

heading(doc, "Ejemplo para defender", 2)
doc.add_paragraph(
    "Al agregar tasks y links no se colocó toda la lógica en la pantalla. El modelo define la forma de los datos, NoteApi los transporta, UpdateNoteHandler valida permisos y contenido, y PrismaNoteRepository persiste. Cada cambio queda localizado en la capa que le corresponde."
)

heading(doc, "8 Patrones utilizados", 1)
table(
    doc,
    ["Patrón", "Uso en el mantenimiento", "Beneficio"],
    [
        ["CQRS", "Commands para crear/actualizar/eliminar y Queries para consultar.", "Casos de uso pequeños y localizables."],
        ["Repository", "Contratos y repositorios Prisma para Note y Binnacle.", "Desacopla negocio y base de datos."],
        ["Facade", "NoteFacade y BinnacleFacade entre páginas, API y Store.", "Simplifica la UI y centraliza loading."],
        ["Mapper", "Conversión Prisma, dominio y respuestas.", "Evita filtrar formatos entre capas."],
        ["DTO", "Validación de entradas HTTP y JSON.", "Rechaza datos incompletos o inválidos."],
        ["Guard", "JWT, roles y guards de rutas Angular.", "Centraliza autenticación y autorización."],
        ["Store reactivo", "Signals para nota, bitácora, loading y error.", "Estado predecible y reutilizable."],
        ["Soft delete", "status igual a cero en lugar de borrado físico.", "Conserva auditoría."],
    ],
    [3.0, 7.0, 6.2],
    7.8,
)

# Quality and defense
heading(doc, "9 Calidad y puntualidad", 1)
heading(doc, "9.1 Puntualidad", 2)
doc.add_paragraph(
    "La entrega vence hoy, 23 de septiembre de 2026, a las 23:59. La defensa debe enviarse antes de esa hora. Para demostrar puntualidad conviene adjuntar este documento, registrar el envío en la plataforma y conservar una captura o comprobante con fecha y hora."
)

heading(doc, "9.2 Evidencias de calidad", 2)
for item in [
    "Compilación Angular completada correctamente.",
    "Backend activo y rutas protegidas por JWT.",
    "Persistencia de tasks y links verificada en Notas.",
    "Validación HTTP/HTTPS en ambos módulos.",
    "Límites de archivos y formatos de imagen.",
    "Estados loading para evitar acciones duplicadas.",
    "Borrado lógico y campos de auditoría.",
    "Separación de vista y edición para reducir errores de uso.",
]:
    bullet(doc, item)

heading(doc, "9.3 Riesgos y siguientes mejoras", 2)
table(
    doc,
    ["Prioridad", "Mejora", "Motivo"],
    [
        ["Alta", "Agregar pruebas automatizadas.", "Proteger validación, permisos y persistencia frente a regresiones."],
        ["Media", "Agregar tareas durante la creación inicial de nota.", "Actualmente se gestionan después desde detalle y edición."],
        ["Media", "Optimizar imágenes grandes antes de guardarlas.", "Reducir memoria, transferencia y almacenamiento."],
        ["Baja", "Reducir el SCSS de create-note.", "Eliminar la advertencia del presupuesto Angular."],
    ],
    [2.2, 6.3, 7.7],
)

heading(doc, "10 Guion para presentar y defender el Sprint 1", 1)
table(
    doc,
    ["Tiempo", "Contenido", "Evidencia"],
    [
        ["1 min", "Alcance real", "Notas: URL y tareas. Mostrar backlog y criterios."],
        ["2 min", "Demostración URL", "Agregar, guardar, recargar y abrir un enlace."],
        ["2 min", "Demostración tareas", "Crear, marcar desde la vista y recargar."],
        ["2 min", "Arquitectura", "Modelo, Facade, API, Handler y repositorio."],
        ["2 min", "SOLID y patrones", "SRP, DIP, CQRS, Repository, Facade y DTO."],
        ["1 min", "Calidad", "Validaciones, permisos, build y persistencia."],
        ["1 min", "Cierre", "Cumplimiento del Sprint 1 y próximos sprints."],
    ],
    [2.0, 6.1, 8.1],
    8,
)

heading(doc, "Demostración paso a paso", 2)
steps = [
    "Iniciar sesión y abrir una nota existente.",
    "Entrar a Editar y agregar una URL relacionada.",
    "Agregar una tarea y guardar la nota.",
    "Volver a la vista y abrir el enlace en una pestaña nueva.",
    "Marcar la tarea sin entrar nuevamente a Editar.",
    "Recargar la página y comprobar que enlace y estado permanecen.",
    "Mostrar note.model.ts, note.api.ts, UpdateNoteHandler y schema.prisma.",
    "Cerrar explicando SOLID y los patrones usados.",
]
for index, item in enumerate(steps, start=1):
    numbered(doc, index, item)

heading(doc, "Mensaje de apertura", 2)
doc.add_paragraph(
    "El mantenimiento asignado se dividió en tres incrementos. En el Sprint 1 trabajamos enlaces URL y tareas en Notas. El resultado conecta la interfaz Angular con casos de uso CQRS y persistencia Prisma, manteniendo validación y permisos. Primero mostraremos el flujo funcional y luego explicaremos las decisiones de diseño."
)

heading(doc, "Mensaje de cierre", 2)
doc.add_paragraph(
    "El Sprint 1 cumple su objetivo: las notas guardan enlaces válidos y tareas que pueden completarse desde la vista. La solución aplica separación de responsabilidades y patrones que permiten continuar con imágenes en el Sprint 2 y reutilizar el enfoque en Bitácora durante el Sprint 3."
)

heading(doc, "11 Preguntas probables y respuestas", 1)
qas = [
    ("¿Qué se pidió exactamente?", "Agregar URL, tareas, imágenes y edición de imágenes en Notas; agregar URL y tareas en Bitácora. El Sprint 1 se concentra en URL y tareas de Notas."),
    ("¿Por qué se dividió así?", "Las listas de enlaces y tareas crean primero el modelo y la persistencia; el editor gráfico tiene mayor complejidad y se separa; Bitácora reutiliza patrones ya probados."),
    ("¿Dónde está SOLID?", "La UI, Facade, Handler y Repository tienen responsabilidades distintas. Los Handlers dependen de contratos de repositorio, lo que aplica inversión de dependencias."),
    ("¿Qué patrones usaron?", "CQRS, Repository, Facade, Mapper, DTO, Guard, Store reactivo y soft delete."),
    ("¿Cómo validan una URL?", "Se normaliza el protocolo y se utiliza new URL; solo se aceptan protocolos HTTP o HTTPS."),
    ("¿Cómo se guardan tareas y enlaces?", "Se envían como JSON dentro de FormData, el backend los analiza y valida, y Prisma los guarda en campos Json."),
    ("¿Cómo se evita editar datos ajenos?", "El JWT entrega el userId; los Handlers comprueban propiedad o membresía. Bitácora también exige el rol estandar."),
    ("¿Qué herramienta se usa para imágenes?", "Fabric.js sobre canvas, con selección, pincel, marcador, recorte, rotación, texto, formas e historial."),
    ("¿Qué falta mejorar?", "Principalmente pruebas automatizadas, optimización de imágenes grandes y permitir tareas en la creación inicial de una nota."),
]
for question, answer in qas:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(question)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    keep(p)
    doc.add_paragraph(answer)

heading(doc, "12 Conclusión", 1)
doc.add_paragraph(
    "El backlog propuesto cubre exactamente el mantenimiento asignado y lo distribuye en tres incrementos defendibles. El Sprint 1 entrega valor visible mediante enlaces y tareas en Notas; el Sprint 2 incorpora imágenes y edición; el Sprint 3 lleva enlaces y tareas a Bitácora. La arquitectura existente permite desarrollar cada incremento sin concentrar responsabilidades ni debilitar los permisos."
)
status(doc, "Sprint 1", GREEN, "URL y tareas de Notas listos para presentación y defensa.")
status(doc, "Sprint 2", TEAL, "imágenes y editor gráfico con Fabric.js.")
status(doc, "Sprint 3", TEAL, "URL y tareas de Bitácora con control por rol y propiedad.")

heading(doc, "Anexo Archivos de referencia", 1)
for item in [
    "prx-frontend/src/app/features/notes",
    "prx-frontend/src/app/features/binnacles",
    "prx-frontend/src/app/core/guards e interceptors",
    "prx-backend/src/modules/notes",
    "prx-backend/src/modules/binnacles",
    "prx-backend/src/modules/auth",
    "prx-backend/prisma/schema.prisma",
]:
    bullet(doc, item, size=9)

for paragraph in doc.paragraphs:
    if paragraph.style and paragraph.style.name.startswith("Heading"):
        keep(paragraph)

for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    if r.font.name is None:
                        r.font.name = "Aptos"
                        r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
                        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")

doc.core_properties.title = "Backlog y defensa del Sprint 1 de mantenimiento PRX"
doc.core_properties.subject = "Notas, Bitácora, imágenes, tareas, URL, SOLID y patrones"
doc.core_properties.author = "Equipo PRX"
doc.save(OUTPUT)
print(OUTPUT)
