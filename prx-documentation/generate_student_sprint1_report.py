from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Informe_Sprint_1_Notas_PRX.docx"
NAVY = "193753"
PALE = "EAF2F7"
TEXT = RGBColor(35, 46, 58)


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), color)


def cell_margins(cell):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for side in ("top", "start", "bottom", "end"):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), "95")
        node.set(qn("w:type"), "dxa")


def row_options(row, repeat=False):
    props = row._tr.get_or_add_trPr()
    props.append(OxmlElement("w:cantSplit"))
    if repeat:
        node = OxmlElement("w:tblHeader")
        node.set(qn("w:val"), "true")
        props.append(node)


def keep_next(paragraph):
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))


def heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    keep_next(paragraph)
    return paragraph


def add_table(doc, headers, rows, widths, font_size=8.6):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    row_options(tbl.rows[0], repeat=True)
    for index, value in enumerate(headers):
        cell = tbl.rows[0].cells[index]
        shade(cell, NAVY)
        cell_margins(cell)
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(value)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(font_size)
    for row_index, values in enumerate(rows):
        row = tbl.add_row()
        row_options(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            if row_index % 2:
                shade(cell, PALE)
            cell_margins(cell)
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            run.font.size = Pt(font_size)
            run.font.color.rgb = TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        r = p.add_run(f"{index}. ")
        r.bold = True
        p.add_run(item)


def code(doc, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    props = p._p.get_or_add_pPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), "F1F4F7")
    props.append(node)
    props.append(OxmlElement("w:keepLines"))
    for line in value.strip().splitlines():
        run = p.add_run(line + "\n")
        run.font.name = "Consolas"
        run.font.size = Pt(8.3)
        run.font.color.rgb = RGBColor(30, 41, 59)


def question(doc, title, answer):
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor(25, 55, 83)
    p = doc.add_paragraph(answer)
    p.paragraph_format.space_after = Pt(8)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = TEXT
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

for style_name, size in (("Title", 29), ("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12)):
    style = doc.styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
    style.paragraph_format.space_after = Pt(7)

title_props = doc.styles["Title"].element.get_or_add_pPr()
title_border = title_props.find(qn("w:pBdr"))
if title_border is not None:
    title_props.remove(title_border)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = header.add_run("PRX  Informe Sprint 1")
run.font.name = "Aptos"
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(91, 105, 120)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Página ").font.size = Pt(8)
page = OxmlElement("w:fldSimple")
page.set(qn("w:instr"), "PAGE")
footer._p.append(page)

# Portada
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(50)
run = p.add_run("PROYECTO PRX")
run.bold = True
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0, 155, 142)

doc.add_paragraph("Informe del Sprint 1", style="Title")
p = doc.add_paragraph()
run = p.add_run("Enlaces URL y tareas en el módulo de Notas")
run.bold = True
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(25, 55, 83)

doc.add_paragraph(
    "En este informe explicamos el trabajo que realizamos durante el Sprint 1 del mantenimiento del proyecto PRX. "
    "Nuestra tarea fue mejorar el módulo de Notas para que el usuario pueda agregar enlaces web y administrar una "
    "lista de tareas. También explicamos los archivos que modificamos, el código principal, los lenguajes utilizados, "
    "la calidad del resultado, los principios SOLID y los patrones de diseño que se aplicaron."
)

add_table(
    doc,
    ["Dato", "Información"],
    [
        ["Actividad", "Presentación y defensa del Sprint 1"],
        ["Fecha límite", "23 de septiembre de 2026 a las 23:59"],
        ["Proyecto", "PRX Mantenimiento evolutivo"],
        ["Módulo trabajado", "Notas"],
        ["Funcionalidades", "Enlaces URL y tareas persistentes"],
        ["Integrantes", "JOSE HUAMANI HUAYPUNA\nVANIA GUARACHI RAMOS\nFERNANDO MARUPA ARNEZ\nCAMILA FATIMA RAMIREZ"],
        ["Evaluación", "Puntualidad  Calidad  SOLID  Patrones"],
    ],
    [1.6, 5.5],
    9,
)

heading(doc, "Resultado del Sprint", 2)
doc.add_paragraph(
    "Al terminar el Sprint, una nota puede guardar enlaces válidos y tareas. Los enlaces se pueden abrir desde "
    "el detalle de la nota y las tareas se pueden marcar como completadas sin entrar nuevamente al formulario de "
    "edición. Los datos quedan guardados en MySQL y se recuperan después de recargar la página."
)

doc.add_page_break()

heading(doc, "1 Resumen de nuestro trabajo", 1)
doc.add_paragraph(
    "Nuestro trabajo consistió en ampliar una función que ya existía. Antes, una nota se concentraba principalmente en "
    "el título, el contenido y los archivos. En este Sprint incorporamos dos listas relacionadas con la nota: enlaces "
    "web y tareas. Para lograrlo tuvimos que realizar cambios tanto en Angular como en NestJS y Prisma."
)

heading(doc, "1.1 Funciones realizadas", 2)
bullets(doc, [
    "Agregar uno o varios enlaces URL a una nota.",
    "Completar automáticamente HTTPS cuando el usuario no escribe el protocolo.",
    "Validar que los enlaces utilicen HTTP o HTTPS.",
    "Mostrar los enlaces guardados y abrirlos en una pestaña nueva.",
    "Agregar y eliminar tareas desde el detalle o la edición de la nota.",
    "Marcar o desmarcar tareas desde el modo de consulta.",
    "Guardar las tareas y enlaces como campos JSON mediante Prisma.",
    "Comprobar los permisos antes de modificar una nota.",
])

heading(doc, "1.2 Alcance del Sprint 1", 2)
doc.add_paragraph(
    "En este Sprint presentamos únicamente los enlaces y las tareas de Notas. Las imágenes y el editor corresponden "
    "al siguiente Sprint, mientras que las mejoras de Bitácora se consideran otro incremento. Esta separación nos "
    "permite defender con claridad lo que realmente corresponde al Sprint 1."
)

heading(doc, "2 Lenguajes y tecnologías utilizadas", 1)
add_table(
    doc,
    ["Lenguaje o tecnología", "Uso en nuestro trabajo", "Motivo"],
    [
        ["TypeScript", "Frontend Angular y backend NestJS", "Permite trabajar con tipos y detectar errores de contrato."],
        ["HTML", "Estructura de la pantalla de detalle", "Muestra enlaces, tareas, botones y checkbox."],
        ["SCSS", "Estilos de las páginas", "Organiza la presentación visual y los estados de los controles."],
        ["Angular 21", "Interfaz de Notas", "Componentes, formularios, rutas y estado reactivo."],
        ["NestJS", "API y casos de uso", "Controladores, inyección de dependencias, DTO y CQRS."],
        ["Prisma", "Persistencia", "Acceso tipado a MySQL y manejo de campos JSON."],
        ["MySQL", "Base de datos", "Almacena las notas, tareas y enlaces."],
        ["JSON", "Formato de las listas", "Transporta y guarda tareas y enlaces relacionados con una nota."],
    ],
    [1.55, 2.55, 3.0],
    8.4,
)

heading(doc, "3 Historias y tareas del Sprint", 1)
add_table(
    doc,
    ["ID", "Historia de usuario", "Puntos", "Estado"],
    [
        ["HU M01", "Como usuario quiero agregar enlaces URL a una nota para relacionarla con recursos externos.", "5", "Hecho"],
        ["HU M02", "Como usuario quiero administrar tareas para controlar actividades pendientes y completadas.", "8", "Hecho"],
    ],
    [0.8, 4.6, 0.7, 1.0],
)

heading(doc, "Tareas técnicas realizadas", 2)
add_table(
    doc,
    ["Tarea", "Trabajo realizado", "Archivo o evidencia"],
    [
        ["T1", "Crear modelos para tareas y enlaces", "note model ts"],
        ["T2", "Agregar controles visuales", "note detail HTML y TS"],
        ["T3", "Normalizar y validar URL", "normalizeUrl e isValidUrl"],
        ["T4", "Enviar las listas al backend", "note api ts"],
        ["T5", "Validar JSON y permisos", "UpdateNoteHandler"],
        ["T6", "Guardar y recuperar listas", "PrismaNoteRepository"],
        ["T7", "Agregar campos a la base de datos", "schema prisma"],
        ["T8", "Comprobar compilación", "ng build y nest build"],
    ],
    [0.65, 3.35, 3.1],
    8.3,
)

heading(doc, "4 Archivos donde están los cambios", 1)
doc.add_paragraph(
    "Estas son las rutas que podemos mostrar durante la defensa. Están escritas desde la carpeta principal del proyecto."
)

heading(doc, "4.1 Frontend", 2)
add_table(
    doc,
    ["Archivo", "Qué contiene"],
    [
        ["prx-frontend/src/app/features/notes/domain/models/note.model.ts", "NoteTaskModel, NoteLinkModel y las listas de la nota."],
        ["prx-frontend/src/app/features/notes/presentation/pages/note-detail-page/note-detail-page.component.html", "Secciones visuales de enlaces y tareas."],
        ["prx-frontend/src/app/features/notes/presentation/pages/note-detail-page/note-detail-page.component.ts", "Agregar, eliminar y marcar tareas; agregar y validar enlaces."],
        ["prx-frontend/src/app/features/notes/domain/requests/update-note.request.ts", "Contrato de actualización con tasks y links."],
        ["prx-frontend/src/app/features/notes/infrastructure/api/note.api.ts", "FormData y envío HTTP al backend."],
    ],
    [4.55, 2.55],
    7.8,
)

heading(doc, "4.2 Backend y base de datos", 2)
add_table(
    doc,
    ["Archivo", "Qué contiene"],
    [
        ["prx-backend/src/modules/notes/domain/entities/note.entity.ts", "NoteTask, NoteLink y contratos del dominio."],
        ["prx-backend/src/modules/notes/application/dto/requests/update-note-request.dto.ts", "Entrada HTTP de tasks y links."],
        ["prx-backend/src/modules/notes/application/commands/update-note/update-note.handler.ts", "Validación, permisos y actualización."],
        ["prx-backend/src/modules/notes/infrastructure/persistence/prisma-note.repository.ts", "Lectura y escritura de JSON."],
        ["prx-backend/src/modules/notes/infrastructure/mappers/note-prisma.mapper.ts", "Conversión de Prisma al dominio."],
        ["prx-backend/prisma/schema.prisma", "Campos tasks y links del modelo Note."],
    ],
    [4.55, 2.55],
    7.8,
)

heading(doc, "5 Código principal que agregamos", 1)
heading(doc, "5.1 Modelos de tareas y enlaces", 2)
code(doc, """
export interface NoteTaskModel {
  id: string;
  title: string;
  completed: boolean;
}

export interface NoteLinkModel {
  id: string;
  url: string;
}
""")
doc.add_paragraph(
    "Este código está en note.model.ts. Lo utilizamos para establecer qué datos debe tener una tarea y qué datos debe "
    "tener un enlace. El identificador permite encontrar cada elemento y completed conserva el estado de la tarea."
)

heading(doc, "5.2 Validación de enlaces", 2)
code(doc, r"""
private normalizeUrl(value: string): string {
  const trimmed = value.trim();
  if (!trimmed) return '';
  return /^https?:\/\//i.test(trimmed) ? trimmed : `https://${trimmed}`;
}

private isValidUrl(value: string): boolean {
  try {
    const url = new URL(value);
    return url.protocol === 'http:' || url.protocol === 'https:';
  } catch {
    return false;
  }
}
""")
doc.add_paragraph(
    "normalizeUrl agrega HTTPS cuando hace falta. isValidUrl comprueba que la dirección tenga una estructura válida "
    "y que el protocolo sea HTTP o HTTPS. El backend vuelve a validar antes de guardar."
)

heading(doc, "5.3 Marcar una tarea desde la vista", 2)
code(doc, """
protected toggleTask(task: NoteTaskModel): void {
  if (!this.canEdit() || this.saving()) return;
  this.tasks.update((tasks) => tasks.map((item) =>
    item.id === task.id ? { ...item, completed: !item.completed } : item
  ));
  if (!this.editing()) this.persist(false);
}
""")
doc.add_paragraph(
    "Este método cambia el estado de una sola tarea. En modo Ver, persist false guarda el cambio sin abrir "
    "el editor. También se comprueba saving para evitar actualizaciones repetidas."
)

heading(doc, "5.4 Envío al backend", 2)
code(doc, """
formData.append('tasks', JSON.stringify(data.tasks));
formData.append('links', JSON.stringify(data.links));
formData.append('retainedFileIds', data.retainedFileIds.join(','));
""")
doc.add_paragraph(
    "Utilizamos FormData porque Notas también trabaja con archivos. Las tareas y los enlaces se convierten a JSON para "
    "enviarlos en la misma petición que los demás datos de la nota."
)

heading(doc, "5.5 Validación y guardado en el backend", 2)
code(doc, """
const tasks = this.parseTasks(command.dto.tasks);
const links = this.parseLinks(command.dto.links);

await this.noteRepository.update(command.id, {
  title: command.dto.title.trim(),
  content: command.dto.content.trim(),
  tasks,
  links,
  updatedBy: command.userId,
});
""")
doc.add_paragraph(
    "UpdateNoteHandler convierte el JSON, valida límites y comprueba permisos. Después llama al repositorio para "
    "guardar la nota. La lógica de negocio no escribe directamente en Prisma."
)

heading(doc, "5.6 Campos de Prisma", 2)
code(doc, """
model Note {
  tasks Json?
  links Json?
}
""")
doc.add_paragraph(
    "Estos campos se encuentran en prx-backend/prisma/schema.prisma. Son opcionales para mantener compatibilidad "
    "con notas antiguas que todavía no tienen tareas o enlaces."
)

heading(doc, "6 Cómo funciona todo el recorrido", 1)
numbered(doc, [
    "El usuario agrega una URL o cambia una tarea en Angular.",
    "El componente valida los datos y actualiza el estado con Signals.",
    "NoteFacade coordina la operación y NoteApi construye FormData.",
    "NotesController recibe la petición y obtiene el usuario desde el JWT.",
    "UpdateNoteHandler verifica permisos y valida las listas.",
    "NoteRepository delega el guardado a PrismaNoteRepository.",
    "Prisma guarda tasks y links como JSON en MySQL.",
    "Los Mappers preparan la respuesta y Angular actualiza la pantalla.",
])

heading(doc, "7 Patrones de diseño utilizados", 1)
add_table(
    doc,
    ["Patrón", "Dónde lo usamos", "Por qué lo usamos"],
    [
        ["CQRS", "Commands para actualizar y Queries para consultar", "Separa operaciones de lectura y escritura."],
        ["Repository", "NoteRepository y PrismaNoteRepository", "Evita que el caso de uso dependa directamente de Prisma."],
        ["Facade", "NoteFacade", "Simplifica la comunicación entre la página, la API y el Store."],
        ["Mapper", "NotePrismaMapper y NoteResponseMapper", "Convierte datos entre Prisma, dominio y respuesta."],
        ["DTO", "UpdateNoteRequestDto", "Define y valida lo que puede recibir el endpoint."],
        ["Store reactivo", "Signals y NoteStore", "Mantiene un estado controlado de datos, carga y errores."],
        ["Guard y Decorator", "JWT y CurrentUser", "Centraliza la identificación y autorización del usuario."],
        ["Soft delete", "Campo status", "Mantiene auditoría sin borrar inmediatamente el registro."],
    ],
    [1.25, 3.15, 2.7],
    8.1,
)

heading(doc, "Explicación para la defensa", 2)
doc.add_paragraph(
    "No utilizamos patrones solamente para mencionarlos. Repository separa la base de datos del caso de uso; Facade "
    "evita que la pantalla coordine todo; Mapper impide mezclar objetos de Prisma con la respuesta; DTO rechaza datos "
    "incorrectos y CQRS mantiene separadas las consultas de las operaciones que modifican información."
)

heading(doc, "8 Principios SOLID", 1)
add_table(
    doc,
    ["Principio", "Aplicación en nuestro trabajo"],
    [
        ["Responsabilidad única", "El componente maneja la interfaz, el Handler las reglas y el Repository los datos."],
        ["Abierto cerrado", "Ampliamos las notas mediante contratos y componentes sin reemplazar toda la arquitectura."],
        ["Sustitución de Liskov", "PrismaNoteRepository puede utilizarse donde se espera el contrato NoteRepository."],
        ["Segregación de interfaces", "NoteApiContract contiene operaciones específicas del módulo de Notas."],
        ["Inversión de dependencias", "UpdateNoteHandler depende de NoteRepository y no directamente de PrismaService."],
    ],
    [2.05, 5.05],
    8.4,
)

doc.add_paragraph(
    "La aplicación más clara de SOLID está en la separación de responsabilidades. No pusimos toda la lógica en el "
    "componente Angular. Cada capa realiza una parte concreta y se comunica mediante contratos."
)

heading(doc, "9 Calidad del trabajo", 1)
heading(doc, "9.1 Evidencias", 2)
bullets(doc, [
    "El frontend Angular compila y genera el bundle correctamente.",
    "El backend NestJS genera Prisma Client y compila sin errores.",
    "La URL se valida en frontend y backend.",
    "Los permisos impiden modificar notas sin acceso al repositorio.",
    "Los enlaces y el estado de las tareas permanecen después de recargar.",
    "Los enlaces externos se abren con noopener y noreferrer.",
    "El estado saving evita enviar la misma operación varias veces.",
])

heading(doc, "9.2 Aspectos que todavía podemos mejorar", 2)
bullets(doc, [
    "Agregar pruebas automatizadas específicas para tareas, enlaces y permisos.",
    "Permitir agregar tareas durante la creación inicial; actualmente se administran desde detalle y edición.",
    "Reducir el archivo SCSS de create note page, que supera el presupuesto Angular por 818 bytes.",
])

heading(doc, "10 Puntualidad", 1)
doc.add_paragraph(
    "La fecha límite indicada fue el 23 de septiembre de 2026 a las 23:59. Para respaldar la puntualidad debemos "
    "adjuntar o conservar el comprobante de entrega de la plataforma. El código y la documentación permiten demostrar "
    "el avance realizado, pero el registro de envío es la evidencia directa del cumplimiento de la fecha."
)

add_table(
    doc,
    ["Evidencia", "Estado"],
    [
        ["Historias de usuario", "Implementadas"],
        ["Compilación frontend", "Correcta"],
        ["Compilación backend", "Correcta"],
        ["Informe de respaldo", "Preparado"],
        ["Comprobante de entrega", "Debe adjuntarse desde la plataforma"],
    ],
    [3.4, 3.7],
)

heading(doc, "Conclusión", 1)
doc.add_paragraph(
    "En este Sprint logramos que el módulo de Notas sea más útil para el seguimiento del trabajo. Los enlaces permiten "
    "relacionar recursos externos y las tareas permiten registrar actividades pendientes o completadas. La solución "
    "no quedó únicamente en la parte visual: incluye validación, permisos, casos de uso y persistencia. Además, la "
    "arquitectura existente permitió aplicar SOLID y patrones sin concentrar toda la lógica en una sola clase."
)

heading(doc, "11 Cómo vamos a presentar", 1)
numbered(doc, [
    "Explicaremos que el Sprint 1 se enfoca en enlaces y tareas de Notas.",
    "Abriremos una nota y agregaremos una URL sin protocolo para mostrar la normalización.",
    "Agregaremos una tarea, guardaremos la nota y volveremos al modo de consulta.",
    "Abriremos el enlace y marcaremos la tarea sin entrar nuevamente a Editar.",
    "Recargaremos la página para demostrar que los datos permanecen.",
    "Mostraremos note-detail-page.component.ts, UpdateNoteHandler y schema.prisma.",
    "Explicaremos Repository, Facade, Mapper, DTO, CQRS y los principios SOLID.",
    "Cerraremos indicando las mejoras pendientes y los siguientes sprints.",
])

heading(doc, "12 Preguntas que podrían hacernos", 1)
question(doc, "¿Qué realizaron en el Sprint 1?", "Agregamos enlaces URL y tareas persistentes al módulo de Notas, incluyendo interfaz, validación, permisos y base de datos.")
question(doc, "¿Dónde se guardan las tareas y los enlaces?", "Se guardan en los campos JSON tasks y links del modelo Note en MySQL mediante Prisma.")
question(doc, "¿Qué lenguajes utilizaron?", "Utilizamos principalmente TypeScript en Angular y NestJS. También utilizamos HTML, SCSS, JSON y el esquema de Prisma para MySQL.")
question(doc, "¿Qué patrones de diseño utilizaron?", "Se utilizaron CQRS, Repository, Facade, Mapper y DTO. Repository es uno de los más importantes porque separa el caso de uso de Prisma.")
question(doc, "¿Dónde aplicaron SOLID?", "En la separación Component, Facade, API, Handler, Repository y Mapper, y en la dependencia del Handler hacia el contrato NoteRepository.")
question(doc, "¿Cómo controlaron los permisos?", "El usuario llega desde el JWT y el Handler comprueba si es propietario o miembro autorizado del repositorio.")
question(doc, "¿Qué falta mejorar?", "Pruebas automatizadas, tareas durante la creación inicial y la reducción del SCSS que supera el presupuesto.")

doc.core_properties.title = "Informe del Sprint 1"
doc.core_properties.subject = "Enlaces URL y tareas en Notas"
doc.core_properties.author = "Proyecto PRX"
doc.save(OUTPUT)
print(OUTPUT)
