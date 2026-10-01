from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Presentacion_y_Defensa_Sprint_1_PRX.docx"
NAVY = "193753"
TEAL = "009B8E"
PALE = "EAF2F7"
TEXT = RGBColor(35, 46, 58)


def fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def margins(cell, top=90, start=100, bottom=90, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def keep_next(paragraph):
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    keep_next(p)
    return p


def table(doc, headers, rows, widths, font_size=8.8):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    header = tbl.rows[0]
    repeat_header(header)
    no_split(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        fill(cell, NAVY)
        margins(cell)
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(font_size)
    for row_index, values in enumerate(rows):
        row = tbl.add_row()
        no_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            if row_index % 2:
                fill(cell, PALE)
            margins(cell)
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
            r.font.color.rgb = TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return tbl


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        r = p.add_run(f"{index}. ")
        r.bold = True
        p.add_run(item)


def label(doc, name, text):
    p = doc.add_paragraph()
    r = p.add_run(name)
    r.bold = True
    r.font.color.rgb = RGBColor(0, 128, 113)
    p.add_run(text)


def code(doc, source):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F1F4F7")
    p_pr.append(shd)
    for line in source.strip().splitlines():
        r = p.add_run(line + "\n")
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(30, 41, 59)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = TEXT
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

for style_name, size in (("Title", 29), ("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12)):
    style = doc.styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
    style.paragraph_format.space_after = Pt(7)

title_p_pr = doc.styles["Title"].element.get_or_add_pPr()
title_border = title_p_pr.find(qn("w:pBdr"))
if title_border is not None:
    title_p_pr.remove(title_border)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run("PRX  Sprint 1")
r.font.name = "Aptos"
r.font.size = Pt(8)
r.font.color.rgb = RGBColor(91, 105, 120)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Página ").font.size = Pt(8)
page = OxmlElement("w:fldSimple")
page.set(qn("w:instr"), "PAGE")
footer._p.append(page)

# Portada
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(52)
r = p.add_run("PROYECTO PRX")
r.bold = True
r.font.size = Pt(12)
r.font.color.rgb = RGBColor(0, 155, 142)

doc.add_paragraph("Presentación y defensa del Sprint 1", style="Title")
p = doc.add_paragraph()
r = p.add_run("Enlaces URL y tareas en Notas")
r.bold = True
r.font.size = Pt(18)
r.font.color.rgb = RGBColor(25, 55, 83)

doc.add_paragraph(
    "Documento de apoyo para presentar y defender el incremento desarrollado durante el Sprint 1. "
    "El alcance se concentra en agregar enlaces web y listas de tareas al módulo de Notas, con "
    "validación, permisos, persistencia y actualización desde la vista."
)

table(
    doc,
    ["Dato", "Detalle"],
    [
        ["Entrega", "Presentación y defensa del Sprint 1"],
        ["Vencimiento", "23 de septiembre de 2026 a las 23:59"],
        ["Módulo", "Notas"],
        ["Incremento", "Enlaces URL y tareas persistentes"],
        ["Criterios", "Puntualidad  Calidad  SOLID  Patrones"],
        ["Estado técnico", "Frontend y backend compilan correctamente"],
    ],
    [1.55, 5.55],
    9,
)

heading(doc, "Conclusión del Sprint", 2)
doc.add_paragraph(
    "El objetivo del Sprint 1 se cumple: una nota puede guardar enlaces válidos, mostrar recursos externos, "
    "administrar tareas y conservar su estado después de recargar. La solución respeta la arquitectura existente "
    "y protege las operaciones mediante autenticación y permisos del repositorio."
)

doc.add_page_break()

heading(doc, "1 Objetivo y alcance", 1)
label(doc, "Objetivo  ", "Convertir una nota en un elemento de seguimiento y consulta, no solamente en texto aislado.")
label(doc, "Valor entregado  ", "El usuario relaciona documentación externa y registra actividades pendientes o completadas dentro de la misma nota.")

heading(doc, "1.1 Incluido en el Sprint 1", 2)
bullets(doc, [
    "Agregar varios enlaces URL a una nota.",
    "Normalizar direcciones sin protocolo y aceptar únicamente HTTP o HTTPS.",
    "Mostrar, abrir y eliminar enlaces relacionados.",
    "Agregar y eliminar tareas desde el detalle o la edición de una nota.",
    "Marcar y desmarcar tareas directamente desde la vista de la nota.",
    "Guardar tareas y enlaces en MySQL mediante Prisma.",
    "Aplicar permisos según propietario, membresía y visibilidad del repositorio.",
])

heading(doc, "1.2 Fuera del Sprint 1", 2)
doc.add_paragraph(
    "Las imágenes y el editor gráfico de Notas corresponden al Sprint 2. Los enlaces y tareas de Bitácora "
    "corresponden al Sprint 3. Pueden mencionarse como continuidad del mantenimiento, pero no deben presentarse "
    "como parte del incremento evaluado en este Sprint 1."
)

heading(doc, "2 Historias de usuario", 1)
table(
    doc,
    ["ID", "Historia", "Puntos", "Estado"],
    [
        ["HU M01", "Como usuario quiero agregar enlaces URL a una nota para relacionarla con recursos externos.", "5", "Hecho"],
        ["HU M02", "Como usuario quiero administrar tareas en una nota para controlar actividades pendientes y completadas.", "8", "Hecho"],
    ],
    [0.75, 4.75, 0.7, 0.9],
)

heading(doc, "2.1 Criterios de aceptación de enlaces", 2)
bullets(doc, [
    "Se puede escribir una dirección con o sin protocolo.",
    "Solo se aceptan enlaces HTTP o HTTPS válidos.",
    "Se pueden agregar varios enlaces y eliminar cualquiera antes de guardar.",
    "Los enlaces guardados aparecen en el detalle y se abren en una pestaña nueva.",
    "Los enlaces permanecen después de recargar la página.",
])

heading(doc, "2.2 Criterios de aceptación de tareas", 2)
bullets(doc, [
    "La tarea debe tener un título no vacío.",
    "Las tareas se pueden agregar y eliminar durante la edición.",
    "El detalle muestra las tareas pendientes y completadas.",
    "El checkbox permite cambiar el estado sin entrar a Editar.",
    "El estado se guarda inmediatamente y permanece después de recargar.",
])

heading(doc, "3 Sprint Backlog terminado", 1)
table(
    doc,
    ["Tarea", "Trabajo técnico", "Evidencia", "Estado"],
    [
        ["T1", "Definir NoteTaskModel y NoteLinkModel", "note model ts", "Hecho"],
        ["T2", "Crear controles para URL", "create note y note detail", "Hecho"],
        ["T3", "Normalizar y validar HTTP o HTTPS", "normalizeUrl e isValidUrl", "Hecho"],
        ["T4", "Crear lista de tareas y checkbox", "note detail HTML y TS", "Hecho"],
        ["T5", "Enviar listas mediante FormData", "note api ts", "Hecho"],
        ["T6", "Analizar y validar JSON en backend", "UpdateNoteHandler", "Hecho"],
        ["T7", "Persistir listas en MySQL", "schema y PrismaNoteRepository", "Hecho"],
        ["T8", "Aplicar permisos del repositorio", "Handlers de Notes", "Hecho"],
        ["T9", "Verificar compilación", "ng build y nest build", "Correcto"],
        ["T10", "Preparar exposición", "Documento y demostración", "Listo"],
    ],
    [0.55, 2.65, 2.95, 0.95],
    8.2,
)

heading(doc, "4 Demostración funcional", 1)
numbered(doc, [
    "Iniciar sesión y abrir una nota existente.",
    "Entrar a Editar y agregar una URL sin protocolo, por ejemplo ejemplo.com.",
    "Comprobar que el sistema la convierte en https://ejemplo.com.",
    "Agregar una tarea con un título breve y guardar la nota.",
    "Regresar al modo Ver y abrir el enlace en una pestaña nueva.",
    "Marcar la tarea desde el checkbox sin entrar nuevamente a Editar.",
    "Recargar la página y comprobar que el enlace y el estado permanecen.",
    "Mostrar un intento con URL inválida para evidenciar la validación.",
])

heading(doc, "Resultado que se debe señalar", 2)
doc.add_paragraph(
    "La demostración prueba el flujo completo: interacción en Angular, envío al backend, validación, autorización, "
    "persistencia en MySQL y recuperación del dato confirmado. No es solamente una maqueta visual."
)

doc.add_page_break()

heading(doc, "5 Código implementado", 1)
heading(doc, "5.1 Contratos de dominio en Angular", 2)
code(doc, """
export interface NoteTaskModel {
  id: string;
  title: string;
  completed: boolean;
}

export interface NoteLinkModel {
  id: string;
  url: string;
}
""")
label(doc, "Qué hace  ", "Define la estructura de los datos que utilizan la pantalla, la API y el Store.")
label(doc, "Por qué se usó  ", "TypeScript detecta errores de contrato antes de ejecutar la aplicación y facilita mantener las capas sincronizadas.")

heading(doc, "5.2 Validación y normalización de URL", 2)
code(doc, r"""
private normalizeUrl(value: string): string {
  const trimmed = value.trim();
  if (!trimmed) return '';
  return /^https?:\/\//i.test(trimmed) ? trimmed : `https://${trimmed}`;
}

private isValidUrl(value: string): boolean {
  try {
    const url = new URL(value);
    return url.protocol === 'http:' || url.protocol === 'https:';
  } catch {
    return false;
  }
}
""")
doc.add_paragraph(
    "normalizeUrl mejora la experiencia al completar HTTPS cuando falta. isValidUrl utiliza el analizador URL del "
    "navegador y limita los protocolos. El backend vuelve a validar porque la interfaz nunca debe ser la única barrera."
)

heading(doc, "5.3 Cambio de estado de tareas", 2)
code(doc, """
protected toggleTask(task: NoteTaskModel): void {
  if (!this.canEdit() || this.saving()) return;
  this.tasks.update((tasks) => tasks.map((item) =>
    item.id === task.id ? { ...item, completed: !item.completed } : item
  ));
  if (!this.editing()) this.persist(false);
}
""")
doc.add_paragraph(
    "La señal se actualiza de forma inmutable. Si el usuario está en Ver, persist false guarda el cambio sin abrir "
    "el formulario de edición. saving evita ejecutar dos actualizaciones al mismo tiempo."
)

heading(doc, "5.4 Transporte mediante FormData", 2)
code(doc, """
formData.append('tasks', JSON.stringify(data.tasks));
formData.append('links', JSON.stringify(data.links));
formData.append('retainedFileIds', data.retainedFileIds.join(','));
""")
doc.add_paragraph(
    "Notas ya utiliza multipart para adjuntos. Por eso las listas se serializan como JSON dentro de FormData, manteniendo "
    "una sola petición compatible con texto y archivos."
)

heading(doc, "5.5 Validación en el caso de uso", 2)
code(doc, """
const tasks = this.parseTasks(command.dto.tasks);
const links = this.parseLinks(command.dto.links);

await this.noteRepository.update(command.id, {
  title: command.dto.title.trim(),
  content: command.dto.content.trim(),
  tasks,
  links,
  updatedBy: command.userId,
});
""")
bullets(doc, [
    "parseTasks exige arreglo, limita a 100 tareas, valida id y título y recorta el título.",
    "parseLinks exige arreglo, limita a 50 enlaces, analiza cada URL y limita a 2048 caracteres.",
    "Antes de actualizar se comprueba que la nota exista y que el usuario tenga permiso.",
])

heading(doc, "5.6 Campos de base de datos", 2)
code(doc, """
model Note {
  tasks Json?
  links Json?
}
""")
doc.add_paragraph(
    "Sí existen campos para guardar correctamente las dos listas. Prisma las escribe como JSON opcional y el Mapper "
    "devuelve arreglos vacíos cuando una nota antigua todavía no tiene esos valores."
)

heading(doc, "6 Arquitectura y recorrido", 1)
table(
    doc,
    ["Capa", "Elemento", "Responsabilidad"],
    [
        ["Presentación", "Note detail component", "Capturar acciones y mostrar tareas y enlaces."],
        ["Aplicación frontend", "NoteFacade", "Coordinar API, loading y Store."],
        ["Infraestructura frontend", "NoteApi", "Construir FormData y ejecutar HTTP."],
        ["Presentación backend", "NotesController", "Recibir la petición y extraer el usuario JWT."],
        ["Aplicación backend", "UpdateNoteHandler", "Validar permisos y reglas del caso de uso."],
        ["Dominio", "NoteEntity y NoteRepository", "Representar datos y contrato de persistencia."],
        ["Infraestructura backend", "PrismaNoteRepository", "Guardar y recuperar JSON en MySQL."],
        ["Conversión", "Prisma Mapper y Response Mapper", "Traducir entre base de datos, dominio y respuesta."],
    ],
    [1.45, 2.1, 3.55],
    8.4,
)

heading(doc, "Flujo completo", 2)
doc.add_paragraph(
    "Usuario → Componente Angular → Facade → NoteApi → NotesController → CommandBus → UpdateNoteHandler → "
    "NoteRepository → Prisma → MySQL → Mapper → respuesta → Store → interfaz."
)

heading(doc, "7 Calidad", 1)
heading(doc, "7.1 Evidencia de calidad", 2)
table(
    doc,
    ["Evidencia", "Resultado", "Qué demuestra"],
    [
        ["Build Angular", "Correcto", "El frontend genera el bundle sin errores."],
        ["Build NestJS", "Correcto", "Prisma Client y TypeScript compilan correctamente."],
        ["Validación URL", "Dos capas", "Frontend mejora UX y backend protege los datos."],
        ["Permisos", "Aplicados", "No se modifican notas sin propiedad o membresía."],
        ["Persistencia", "Comprobable", "Los datos sobreviven a una recarga."],
        ["Enlaces externos", "Protegidos", "Se usa noopener noreferrer al abrir otra pestaña."],
        ["Estado saving", "Aplicado", "Evita acciones duplicadas durante una petición."],
        ["Borrado lógico", "Aplicado", "Se mantiene trazabilidad mediante status."],
    ],
    [1.45, 1.45, 4.2],
    8.5,
)

heading(doc, "7.2 Definición de terminado", 2)
bullets(doc, [
    "La interfaz administra enlaces y tareas sin errores de maquetación.",
    "Frontend y backend compilan.",
    "Los datos se guardan y se recuperan mediante Prisma.",
    "Las entradas inválidas reciben mensajes comprensibles.",
    "Los permisos impiden modificar notas sin acceso al repositorio.",
    "La demostración puede repetirse antes de la entrega.",
])

heading(doc, "7.3 Mejoras pendientes", 2)
bullets(doc, [
    "Agregar pruebas automatizadas unitarias y de integración para las nuevas funciones.",
    "Permitir agregar tareas directamente durante la creación inicial de una nota; actualmente se administran desde detalle y edición.",
    "Reducir el SCSS de create note page, que supera el presupuesto Angular por 818 bytes.",
])

heading(doc, "8 Principios SOLID", 1)
table(
    doc,
    ["Principio", "Aplicación en el Sprint 1", "Defensa"],
    [
        ["S Responsabilidad única", "Component, Facade, API, Handler, Repository y Mapper tienen tareas distintas.", "Los cambios quedan localizados en la capa responsable."],
        ["O Abierto cerrado", "Contratos y componentes permiten ampliar comportamientos sin reescribir consumidores.", "Se agregaron listas respetando la arquitectura existente."],
        ["L Sustitución de Liskov", "PrismaNoteRepository cumple el contrato NoteRepository.", "El caso de uso puede trabajar con cualquier implementación válida."],
        ["I Segregación de interfaces", "NoteApiContract contiene operaciones específicas de Notas.", "El módulo no depende de métodos ajenos a su necesidad."],
        ["D Inversión de dependencias", "UpdateNoteHandler inyecta NoteRepository abstracto.", "La lógica de aplicación no depende directamente de Prisma."],
    ],
    [1.35, 3.55, 2.2],
    8.0,
)

heading(doc, "Cómo explicarlo oralmente", 2)
doc.add_paragraph(
    "No colocamos toda la lógica en la pantalla. El componente maneja interacción; la Facade coordina; la API "
    "transporta; el Handler aplica reglas y permisos; el Repository persiste; el Mapper transforma datos. Esta división "
    "es la evidencia principal de responsabilidad única e inversión de dependencias."
)

heading(doc, "9 Patrones de diseño", 1)
table(
    doc,
    ["Patrón", "Uso en el Sprint 1", "Por qué se usó"],
    [
        ["CQRS", "UpdateNoteCommand para escribir y Queries para consultar.", "Separar lecturas y cambios en casos de uso pequeños."],
        ["Repository", "NoteRepository y PrismaNoteRepository.", "Desacoplar negocio de Prisma y MySQL."],
        ["Facade", "NoteFacade entre las páginas, API y Store.", "Simplificar el componente y centralizar coordinación."],
        ["Mapper", "NotePrismaMapper y NoteResponseMapper.", "Evitar filtrar formatos de persistencia a la interfaz."],
        ["DTO", "UpdateNoteRequestDto.", "Definir y validar la entrada HTTP."],
        ["Reactive Store", "Signals y NoteStore.", "Mantener estado predecible de datos, carga y error."],
        ["Guard y Decorator", "JWT y CurrentUser.", "Centralizar identidad y autorización."],
        ["Soft delete", "status igual a cero.", "Conservar auditoría sin eliminar físicamente."],
    ],
    [1.3, 3.25, 2.55],
    8.1,
)

heading(doc, "Cómo defender los patrones", 2)
doc.add_paragraph(
    "Los patrones no se usaron solo para nombrarlos. CQRS organiza operaciones, Repository aísla la base de datos, "
    "Facade reduce dependencias de la UI, Mapper protege límites entre capas y DTO impide que datos incorrectos "
    "lleguen al caso de uso. Cada patrón responde a un problema concreto del Sprint."
)

heading(doc, "10 Puntualidad", 1)
doc.add_paragraph(
    "La entrega vence el 23 de septiembre de 2026 a las 23:59. Para demostrar puntualidad se debe enviar el trabajo "
    "antes de esa hora y conservar evidencia del envío. El código y el documento están preparados; la confirmación final "
    "en la plataforma corresponde al estudiante."
)
table(
    doc,
    ["Evidencia", "Estado", "Acción antes de entregar"],
    [
        ["Historias del Sprint", "Hechas", "Demostrar ambos flujos."],
        ["Frontend", "Compila", "Conservar salida de ng build."],
        ["Backend", "Compila", "Conservar salida de nest build."],
        ["Documento", "Preparado", "Adjuntarlo con la entrega."],
        ["Registro de envío", "Pendiente del estudiante", "Guardar captura con fecha y hora."],
    ],
    [1.7, 1.7, 3.7],
)

heading(doc, "11 Guion de exposición", 1)
table(
    doc,
    ["Tiempo", "Tema", "Qué decir o mostrar"],
    [
        ["1 min", "Objetivo", "Explicar el problema y el alcance real del Sprint 1."],
        ["3 min", "Demostración", "URL, tarea, guardado, checkbox y recarga."],
        ["2 min", "Código", "Modelo, validación, FormData, Handler y Prisma."],
        ["2 min", "Calidad", "Builds, validación en dos capas, permisos y límites."],
        ["2 min", "SOLID", "SRP y DIP con la separación de capas."],
        ["2 min", "Patrones", "CQRS, Repository, Facade, Mapper y DTO."],
        ["1 min", "Cierre", "Valor entregado, alcance cumplido y siguientes sprints."],
    ],
    [0.8, 1.4, 4.9],
)

heading(doc, "11.1 Mensaje de apertura", 2)
doc.add_paragraph(
    "En el Sprint 1 trabajamos el módulo de Notas. El objetivo fue permitir que una nota relacionara recursos web y "
    "administrara tareas pendientes o completadas. La solución abarca interfaz, validación, permisos, backend y persistencia."
)

heading(doc, "11.2 Mensaje de cierre", 2)
doc.add_paragraph(
    "El incremento cumple las dos historias comprometidas. Los enlaces y tareas se guardan, se recuperan y respetan "
    "los permisos del repositorio. La arquitectura separada permite continuar en el Sprint 2 con imágenes y en el "
    "Sprint 3 con Bitácora sin concentrar toda la lógica en un solo componente."
)

heading(doc, "12 Preguntas probables", 1)
qa = [
    ("¿Qué entregaron exactamente?", "Enlaces URL y tareas persistentes en Notas, incluyendo vista, edición, validación y permisos."),
    ("¿Dónde se guardan?", "En los campos JSON opcionales tasks y links del modelo Note en MySQL mediante Prisma."),
    ("¿Cómo validan una URL?", "Angular normaliza y utiliza new URL; el backend vuelve a analizar el valor antes de guardarlo."),
    ("¿Por qué usaron JSON?", "Las listas son pequeñas, pertenecen a una sola nota y se recuperan junto con ella. Si luego necesitan consultas globales, podrían normalizarse en tablas."),
    ("¿Cómo evitan editar notas ajenas?", "El JWT entrega el userId y el Handler comprueba propiedad o membresía según la visibilidad del repositorio."),
    ("¿Qué pasa al marcar una tarea desde Ver?", "Signals actualiza la lista y persist false ejecuta PUT sin abrir el editor. El cambio queda guardado."),
    ("¿Dónde está SOLID?", "En la separación Component, Facade, API, Handler, Repository y Mapper, y en la dependencia del Handler hacia un contrato abstracto."),
    ("¿Qué patrones usaron?", "CQRS, Repository, Facade, Mapper, DTO, Store reactivo, Guards y soft delete."),
    ("¿Cómo demuestran calidad?", "Con compilaciones correctas, validación en dos capas, permisos, límites y persistencia comprobable."),
    ("¿Qué falta mejorar?", "Pruebas automatizadas, tareas en la creación inicial y reducir el SCSS que supera el presupuesto."),
]
for question, answer in qa:
    p = doc.add_paragraph()
    r = p.add_run(question)
    r.bold = True
    r.font.color.rgb = RGBColor(25, 55, 83)
    p = doc.add_paragraph(answer)
    p.paragraph_format.space_after = Pt(8)

heading(doc, "13 Lista final antes de enviar", 1)
bullets(doc, [
    "Abrir el documento y revisar nombre, fecha y alcance.",
    "Ensayar la demostración con una nota de prueba.",
    "Tener frontend, backend y base de datos listos antes de presentar.",
    "No incluir imágenes ni Bitácora como trabajo del Sprint 1.",
    "Enviar antes de las 23:59 y conservar captura del comprobante.",
])

doc.core_properties.title = "Presentación y defensa del Sprint 1"
doc.core_properties.subject = "Enlaces URL y tareas en Notas"
doc.core_properties.author = "Proyecto PRX"
doc.save(OUTPUT)
print(OUTPUT)
